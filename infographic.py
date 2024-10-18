
import os
from pptx import Presentation
from fpdf import FPDF

def create_infographic(summary_data, output_file):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Add text to the PDF
    pdf.cell(200, 10, "Client Name: {}".format(summary_data.get('ClientName', 'N/A')), ln=True)
    pdf.cell(200, 10, "Financial Well Being: {}".format(summary_data.get('FinancialWellBeing', 'N/A')), ln=True)
    pdf.cell(200, 10, "Key Wins: {}".format(summary_data.get('KeyWins', 'N/A')), ln=True)
    pdf.cell(200, 10, "Goals: {}".format(summary_data.get('Goals', 'N/A')), ln=True)
    pdf.cell(200, 10, "Estate Planning: {}".format(summary_data.get('EstatePlanning', 'N/A')), ln=True)
    pdf.cell(200, 10, "Your Tax Rates: {}".format(summary_data.get('YourTaxRates', 'N/A')), ln=True)
    pdf.cell(200, 10, "Net Worth: {}".format(summary_data.get('NetWorth', 'N/A')), ln=True)
    pdf.cell(200, 10, "Asset Allocation: {}".format(summary_data.get('AssetAllocation', 'N/A')), ln=True)
    pdf.cell(200, 10, "Investment Tax Allocation: {}".format(summary_data.get('InvestmentTaxAllocation', 'N/A')), ln=True)
    pdf.cell(200, 10, "Homework: {}".format(summary_data.get('Homework', 'N/A')), ln=True)

    # Save the PDF
    pdf.output(output_file)